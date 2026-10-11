"""Tag an untagged PowerPoint deck so the app knows which shapes to refresh.

The refresher (pptx_refresher.py) only touches shapes whose *name* starts with `kpi:` or `chart:`.
This tool applies those names to a deck from a small JSON spec, so tagging is repeatable instead of a
manual Selection-Pane exercise:

    python template_tagger.py untagged.pptx tagged.pptx --spec spec.json

Spec format (slide numbers are 1-based):

    {
      "chart_tags": {"2": ["academy_learners_by_country", "academy_learners_by_language"]},
      "kpi_tags":   {"2": [["508", "academy_total_learners"], ["17", "academy_countries"]]}
    }

* chart_tags — picture keys for a slide, assigned to the slide's large pictures from left to right
  (small decorative logos are ignored).
* kpi_tags   — [visible text, key] pairs. The shape whose text equals the visible text exactly gets the
  name `kpi:<key>`; repeated texts are matched in shape order.

You can inspect or edit the names in PowerPoint via Home > Arrange > Selection Pane.
A ready-made demo spec lives in demo_data/demo_template_spec.json (see demo_data/build_demo_assets.py).
"""

from __future__ import annotations

import argparse
import io
import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# Pictures smaller than this (EMU, 1 inch) in both directions are treated as decorative logos.
MIN_CHART_EMU = 914400


def clean(text: str) -> str:
    return " ".join(text.replace("\xa0", " ").split())


def walk_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_shapes(shape.shapes)


def get_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return clean("\n".join(p.text for p in shape.text_frame.paragraphs))


def load_spec(path: Path) -> tuple[dict[int, list[str]], dict[int, list[tuple[str, str]]]]:
    """Read a tagging spec; returns (chart_tags, kpi_tags) keyed by 1-based slide number."""
    raw = json.loads(Path(path).read_text(encoding="utf-8"))
    charts = {int(k): list(v) for k, v in raw.get("chart_tags", {}).items()}
    kpis = {int(k): [(str(t), str(key)) for t, key in v] for k, v in raw.get("kpi_tags", {}).items()}
    return charts, kpis


def tag_charts(slide, tags: list[str]) -> int:
    if not tags:
        return 0
    pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    pictures = [p for p in pictures if p.width > MIN_CHART_EMU or p.height > MIN_CHART_EMU]
    pictures.sort(key=lambda p: (p.left, p.top))
    count = 0
    for shape, tag in zip(pictures, tags):
        shape.name = f"chart:{tag}"
        count += 1
    return count


def tag_kpis(slide, targets: list[tuple[str, str]]) -> int:
    if not targets:
        return 0
    shapes = list(walk_shapes(slide.shapes))
    used: set[int] = set()
    count = 0
    for visible_text, key in targets:
        wanted = clean(visible_text)
        for idx, shape in enumerate(shapes):
            if idx not in used and get_text(shape) == wanted:
                shape.name = f"kpi:{key}"
                used.add(idx)
                count += 1
                break
    return count


def parse_spec(raw: dict) -> tuple[dict[int, list[str]], dict[int, list[tuple[str, str]]]]:
    """Same as load_spec but from an already-parsed dict (used by the UI)."""
    charts = {int(k): list(v) for k, v in raw.get("chart_tags", {}).items()}
    kpis = {int(k): [(str(t), str(key)) for t, key in v] for k, v in raw.get("kpi_tags", {}).items()}
    return charts, kpis


def inventory(prs) -> list[dict]:
    """Every taggable shape per slide (text boxes and large pictures), to help write a spec."""
    rows = []
    for number, slide in enumerate(prs.slides, 1):
        for shape in walk_shapes(slide.shapes):
            text = get_text(shape)
            is_chart = (shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                        and (shape.width > MIN_CHART_EMU or shape.height > MIN_CHART_EMU))
            if text or is_chart:
                rows.append({"slide": number, "kind": "picture" if is_chart else "text",
                             "current name": shape.name, "text": text[:60]})
    return rows


def tag_bytes(deck: bytes, spec: dict) -> tuple[bytes, int, int, list[str]]:
    """Tag a deck held in memory. Returns (new deck, kpi count, chart count, unmatched spec entries)."""
    chart_tags, kpi_tags = parse_spec(spec)
    prs = Presentation(io.BytesIO(deck))
    n_chart = n_kpi = 0
    unmatched: list[str] = []
    for idx, slide in enumerate(prs.slides, 1):
        got = tag_charts(slide, chart_tags.get(idx, []))
        n_chart += got
        if got < len(chart_tags.get(idx, [])):
            unmatched.append(f"slide {idx}: {len(chart_tags[idx]) - got} chart tag(s) without a picture")
        before = n_kpi
        n_kpi += tag_kpis(slide, kpi_tags.get(idx, []))
        if n_kpi - before < len(kpi_tags.get(idx, [])):
            unmatched.append(f"slide {idx}: {len(kpi_tags[idx]) - (n_kpi - before)} KPI text(s) not found")
    for idx in (set(chart_tags) | set(kpi_tags)) - set(range(1, len(prs.slides) + 1)):
        unmatched.append(f"slide {idx} does not exist in this deck")
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue(), n_kpi, n_chart, unmatched


def tag_template(input_path: Path, output_path: Path, spec_path: Path) -> tuple[int, int]:
    """Tag `input_path` per the spec and save to `output_path`. Returns (kpi_count, chart_count)."""
    chart_tags, kpi_tags = load_spec(spec_path)
    prs = Presentation(input_path)
    chart_count = kpi_count = 0
    for idx, slide in enumerate(prs.slides, 1):
        chart_count += tag_charts(slide, chart_tags.get(idx, []))
        kpi_count += tag_kpis(slide, kpi_tags.get(idx, []))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    return kpi_count, chart_count


def main() -> int:
    parser = argparse.ArgumentParser(description="Tag a PowerPoint deck's KPI text boxes and chart pictures.")
    parser.add_argument("input", type=Path, help="untagged .pptx")
    parser.add_argument("output", type=Path, help="where to write the tagged .pptx")
    parser.add_argument("--spec", type=Path, required=True, help="JSON tagging spec")
    args = parser.parse_args()
    kpis, charts = tag_template(args.input, args.output, args.spec)
    print(f"Tagged {kpis} KPI text boxes and {charts} chart/image placeholders")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
