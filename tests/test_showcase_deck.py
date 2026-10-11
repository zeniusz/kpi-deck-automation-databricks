"""The anonymised showcase deck: every tag is known to the app, and nothing from a branded template is left."""

import re
import zipfile
from pathlib import Path

from pptx import Presentation

import kpis

DECK = Path(__file__).resolve().parent.parent / "templates" / "Analytics_KPIs_showcase_tagged.pptx"


def _walk(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == 6:
            yield from _walk(shape.shapes)


def test_every_tag_is_a_known_kpi_or_chart_key():
    deck = Presentation(DECK)
    names = {s.name for slide in deck.slides for s in _walk(slide.shapes) if s.name.startswith(("kpi:", "chart:"))}
    scalars, charts = {i["key"] for i in kpis.SCALARS}, {i["key"] for i in kpis.CHARTS}
    unknown = [n for n in names if (n.startswith("kpi:") and n[4:] not in scalars)
               or (n.startswith("chart:") and n[6:] not in charts)]
    assert len(names) > 60 and not unknown


def test_package_is_clean():
    z = zipfile.ZipFile(DECK)
    names = z.namelist()
    assert not [n for n in names if "embeddings" in n or "changesInfos" in n or n.endswith("authors.xml")]
    assert len([n for n in names if n.startswith("ppt/slideLayouts/slideLayout")]) <= 12
    assert len([n for n in names if n.startswith("ppt/slideMasters/slideMaster")]) == 1
    text = " ".join(z.read(n).decode("utf8", "ignore") for n in names if n.endswith((".xml", ".rels")))
    assert not re.search(r"hlinkClick|buBlip|custDataLst", text)
    assert "Demo" in z.read("docProps/core.xml").decode()
