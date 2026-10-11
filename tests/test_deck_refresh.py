"""Template-tagging mechanism: tagging a deck from a spec, then refreshing tagged shapes by name."""

import io
import json
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

import kpis
import pptx_refresher as pr
import template_tagger as tt
from ui_theme import deck_tags

ROOT = Path(__file__).resolve().parent.parent
UNTAGGED = ROOT / "demo_data" / "Analytics_KPIs_demo_untagged.pptx"
TAGGED = ROOT / "templates" / "Analytics_KPIs_demo_template_tagged.pptx"
SPEC = ROOT / "demo_data" / "demo_template_spec.json"


def _shape_names(path_or_bytes):
    source = io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes) else path_or_bytes
    deck = Presentation(source)
    return {(i, s.name) for i, slide in enumerate(deck.slides, 1) for s in slide.shapes
            if s.name.startswith(("kpi:", "chart:"))}


@pytest.fixture
def chart_pngs(tmp_path):
    images = {}
    for chart in kpis.CHARTS:
        path = tmp_path / f"{chart['key']}.png"
        Image.new("RGB", (800, 400), (200, 60, 60)).save(path)
        images[chart["key"]] = path
    return images


def test_tagger_reproduces_the_shipped_template(tmp_path):
    out = tmp_path / "tagged.pptx"
    n_kpi, n_chart = tt.tag_template(UNTAGGED, out, SPEC)
    assert (n_kpi, n_chart) == (14, 8)
    assert _shape_names(out) == _shape_names(TAGGED)


def test_tagger_ignores_spec_entries_that_do_not_match(tmp_path):
    spec = tmp_path / "spec.json"
    spec.write_text(json.dumps({"kpi_tags": {"2": [["no such text", "ghost"]]}, "chart_tags": {}}))
    kpi_count, chart_count = tt.tag_template(UNTAGGED, tmp_path / "o.pptx", spec)
    assert (kpi_count, chart_count) == (0, 0)


def test_deck_tags_excludes_manual_slides():
    keys = deck_tags(TAGGED.read_bytes(), "kpi")
    assert "academy_total_learners" in keys
    assert "group_employees" not in keys  # tagged, but on a manual slide


def test_refresh_replaces_tagged_values_only(chart_pngs):
    values = {"academy_total_learners": 1234, "report_month_label": "July 2026 Data",
              "db_training_coverage": 0.42, "group_employees": 999}
    refreshed, stats = pr.refresh_powerpoint(
        TAGGED.read_bytes(), scalar_config=kpis.SCALARS, scalar_values=values, chart_images=chart_pngs)
    deck = Presentation(io.BytesIO(refreshed))
    by_name = {s.name: s for slide in deck.slides for s in slide.shapes}

    assert by_name["kpi:academy_total_learners"].text_frame.text == "1,234"
    assert by_name["kpi:report_month_label"].text_frame.text == "July 2026 Data"
    assert by_name["kpi:db_training_coverage"].text_frame.text == "0.42%"
    # untagged text is left exactly as designed
    assert any(s.has_text_frame and s.text_frame.text == "Total learners"
               for s in deck.slides[1].shapes)
    # KPIs without a supplied value keep their template text
    assert by_name["kpi:academy_countries"].text_frame.text == "5"
    assert stats["kpi_updates"] == 3 and stats["chart_updates"] == 8


def test_manual_slide_is_never_refreshed(chart_pngs):
    assert 6 in pr.MANUAL_SLIDE_NUMBERS
    refreshed, _ = pr.refresh_powerpoint(
        TAGGED.read_bytes(), scalar_config=kpis.SCALARS,
        scalar_values={"group_employees": 999}, chart_images=chart_pngs)
    manual = Presentation(io.BytesIO(refreshed)).slides[5]
    texts = [s.text_frame.text for s in manual.shapes if s.has_text_frame]
    assert "Edited by hand" in texts and "999" not in texts


def test_picture_keeps_its_tag_so_next_month_can_reuse_the_deck(chart_pngs):
    once, _ = pr.refresh_powerpoint(TAGGED.read_bytes(), scalar_config=kpis.SCALARS,
                                    scalar_values={}, chart_images=chart_pngs)
    assert {n for _, n in _shape_names(once)} >= {"chart:academy_learners_by_country"}
    twice, stats = pr.refresh_powerpoint(once, scalar_config=kpis.SCALARS,
                                         scalar_values={}, chart_images=chart_pngs)
    assert stats["chart_updates"] == 8


def test_chart_box_override_wins_over_placeholder_bounds(chart_pngs):
    left, top, width, height = pr.CHART_BOX_OVERRIDES["community_geographic_footprint"]
    refreshed, _ = pr.refresh_powerpoint(TAGGED.read_bytes(), scalar_config=kpis.SCALARS,
                                         scalar_values={}, chart_images=chart_pngs)
    pic = next(s for s in Presentation(io.BytesIO(refreshed)).slides[3].shapes
               if s.name == "chart:community_geographic_footprint")
    emu = 914400
    assert left * emu <= pic.left and pic.left + pic.width <= (left + width) * emu + 1
    assert top * emu <= pic.top and pic.top + pic.height <= (top + height) * emu + 1


def test_manual_chart_keys_are_skipped(chart_pngs, monkeypatch):
    monkeypatch.setattr(pr, "MANUAL_CHART_KEYS", {"academy_learners_by_country"})
    refreshed, stats = pr.refresh_powerpoint(TAGGED.read_bytes(), scalar_config=kpis.SCALARS,
                                             scalar_values={}, chart_images=chart_pngs)
    assert stats["chart_updates"] == 7


def test_month_labels_roll_forward_together():
    assert pr.build_month_replacements("2026.06", "2026.07") == {
        "2026.06": "2026.07", "2026-06": "2026-07", "June 2026": "July 2026"}


@pytest.mark.parametrize("value,fmt,expected", [
    (1234, "auto", "1,234"), (12.5, "auto", "12.5"), (1020, "compact", "1.02K"),
    (0.4, "percent", "0.4%"), (None, "auto", ""), ("x", "text", "x"),
])
def test_format_value(value, fmt, expected):
    assert pr.format_value(value, fmt) == expected


def test_every_dashboard_tile_targets_a_known_dashboard_and_chart():
    chart_keys = {c["key"] for c in kpis.CHARTS}
    for key, spec in kpis.DASHBOARD_TILES.items():
        assert spec["dashboard"] in kpis.DATABRICKS_DASHBOARDS, key
        assert key in chart_keys, f"{key} has a tile mapping but no chart definition"


def test_every_scalar_and_chart_key_is_unique():
    for items in (kpis.SCALARS, kpis.CHARTS):
        keys = [i["key"] for i in items]
        assert len(keys) == len(set(keys))


def test_tag_bytes_reports_unmatched_spec_entries():
    spec = {"chart_tags": {"2": ["a", "b", "c", "d", "e"]}, "kpi_tags": {"2": [["no such text", "ghost"]], "99": [["x", "y"]]}}
    _, n_kpi, n_chart, unmatched = tt.tag_bytes(UNTAGGED.read_bytes(), spec)
    assert n_kpi == 0 and n_chart >= 1
    assert any("KPI text" in u for u in unmatched) and any("slide 99" in u for u in unmatched)


def test_tag_bytes_matches_file_based_tagger():
    spec = json.loads(SPEC.read_text(encoding="utf-8"))
    tagged, n_kpi, n_chart, unmatched = tt.tag_bytes(UNTAGGED.read_bytes(), spec)
    assert (n_kpi, n_chart, unmatched) == (14, 8, [])
    assert _shape_names(tagged) == _shape_names(TAGGED)
