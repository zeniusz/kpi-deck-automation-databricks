"""Generate the fictional demo assets: synthetic Excel exports, a demo deck, its tagging spec and the
tagged template.

Everything is invented and deterministic (fixed seed). Names are common given/family names combined at
random; e-mail addresses use the reserved `.example` domains from config/org_profile.toml.

    python demo_data/build_demo_assets.py

Writes:
    demo_data/excel/*.xlsx                      synthetic uploads, one per routing rule
    demo_data/demo_template_spec.json           tagging spec for the demo deck
    demo_data/Analytics_KPIs_demo_untagged.pptx the demo deck before tagging
    templates/Analytics_KPIs_demo_template_tagged.pptx   the same deck after template_tagger.py
"""

from __future__ import annotations

import io
import json
import random
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import org_profile  # noqa: E402
import template_tagger  # noqa: E402

RNG = random.Random(20260710)
FIRST = ["Anna", "Ben", "Clara", "David", "Eva", "Felix", "Greta", "Hans", "Ida", "Jonas", "Klara",
         "Lukas", "Mia", "Nico", "Olga", "Paul", "Rita", "Sven", "Tina", "Urs", "Vera", "Wim"]
LAST = ["Schmidt", "Meier", "Novak", "Kowalski", "Garcia", "Silva", "de Vries", "Keller", "Weber",
        "Fischer", "Dvorak", "Nowak", "Bauer", "Costa", "Jansen", "Lehner", "Moser", "Pereira"]
DOMAINS = [d for d in org_profile.DOMAIN_COUNTRY if "contractor" not in d and "legacy" not in d]
EXCEL = ROOT / "demo_data" / "excel"


def person() -> tuple[str, str]:
    return RNG.choice(FIRST), RNG.choice(LAST)


def email(first: str, last: str, domain: str | None = None) -> str:
    domain = domain or RNG.choice(DOMAINS)
    return f"{first}.{last}".lower().replace(" ", "") + "@" + domain


def write(name: str, frame: pd.DataFrame) -> None:
    EXCEL.mkdir(parents=True, exist_ok=True)
    frame.to_excel(EXCEL / name, index=False)


def build_excel() -> None:
    # 1. Teams listing snapshot → APPEND (Country/Company derived from the e-mail domain)
    rows = []
    for _ in range(60):
        f, l = person()
        rows.append({"User": email(f, l), "Name": f"{f} {l}", "Company": None, "Country": None})
    for _ in range(3):
        f, l = person()
        rows.append({"User": None, "Name": f"{f} {l}", "Company": None, "Country": None})
    write("Analytics Community 2026.07.xlsx", pd.DataFrame(rows))

    # 2. Registration Forms export → FULL_REPLACE (the real address is in Email2)
    iso = list(org_profile.COUNTRY_BY_ISO)
    rows = []
    for i in range(1, 41):
        f, l = person()
        code = RNG.choice(iso)
        rows.append({
            "ID": i, "Start time": f"2026-06-{RNG.randint(1, 28):02d}",
            "Completion time": f"2026-06-{RNG.randint(1, 28):02d}",
            "Email": "anonymous", "Email2": email(f, l),
            "First name": f, "Last name": l,
            "Company": f"{RNG.choice(['Northwind', 'Fabrikam', 'Tailspin', 'Woodgrove'])} ({code})",
            "Department": RNG.choice(["Finance", "Operations", "IT", "Sales"]),
            "Job title": RNG.choice(["Analyst", "Manager", "Engineer", "Consultant"]),
            "What contents and formats are you especially interested in?":
                RNG.choice(["Use cases", "Python and SQL", "Governance and ethics", "Dashboards"]),
            "Country": None,
        })
    rows.append(dict(rows[3]))  # a repeat submission, to show de-duplication
    write("Analytics Community (2026.07.03).xlsx", pd.DataFrame(rows))

    # 3. Lakehouse Guild membership snapshot → APPEND (separate from the community listing)
    rows = []
    for _ in range(25):
        f, l = person()
        rows.append({"DisplayName": f"{f} {l}", "Email": email(f, l), "Roles": "member",
                     "Company": None, "Country": None})
    write("Lakehouse Guild TEAMS listing 2026.07.xlsx", pd.DataFrame(rows))

    # 4. Learning course export (Atlas platform) → NEW_TABLE_PERIOD; no e-mail column
    rows = [{"Company": f"{RNG.choice(['Northwind', 'Fabrikam', 'Tailspin'])} ({RNG.choice(iso)})",
             "Status": RNG.choice(["Passed", "In progress", "Enrolled"]), "Country": None}
            for _ in range(30)]
    write("Atlas_DataGovernance_Basics.xlsx", pd.DataFrame(rows))

    # 5. Raw LMS export: the numeric id prefix decides the platform (fictional id, see learning_metadata)
    rows = [{"Company": f"Woodgrove ({RNG.choice(iso)})",
             "Status": RNG.choice(["Passed", "In progress"]), "Country": None} for _ in range(20)]
    write("300400_export_eLearningDataBasicsAwarenessPart1_20260803_0838.xlsx", pd.DataFrame(rows))

    # 6. Event attendance → NEW_TABLE_EVENT (≤1 minute attendees are dropped, repeats de-duplicated)
    rows = []
    for _ in range(30):
        f, l = person()
        rows.append({"Name": f"{f} {l}", "Email": email(f, l),
                     "Dauer": f"{RNG.randint(0, 1)}h {RNG.randint(0, 59)}m {RNG.randint(0, 59)}s"})
    rows.append({"Name": "Brief Visitor", "Email": email("brief", "visitor"), "Dauer": "0m 40s"})
    rows.append(dict(rows[0], Dauer="0h 5m 3s"))
    write("Analytics Event Number 5.xlsx", pd.DataFrame(rows))

    # 7. Revenue demo table → APPEND with a reporting-month tag
    rows = [{"Company": f"Company {i:02d}", "Country": RNG.choice(list(org_profile.COUNTRY_BY_ISO.values())),
             "Annual Revenue (EUR)": RNG.randint(20, 900) * 1_000_000} for i in range(1, 41)]
    write("Revenue_2026.07.xlsx", pd.DataFrame(rows))


# ── Demo deck ────────────────────────────────────────────────────────────────────────────────────
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
ACCENT = RGBColor(0x2D, 0x9C, 0xDB)
COUNTRIES = ["Germany", "Austria", "Poland", "Spain", "Switzerland"]


def chart_png(kind: str, seed: int) -> bytes:
    rng = random.Random(seed)
    fig, ax = plt.subplots(figsize=(6, 3.4), dpi=110)
    if kind == "donut":
        vals = [rng.randint(30, 90), rng.randint(5, 30)]
        ax.pie(vals, labels=["English", "German"], wedgeprops={"width": 0.45},
               colors=["#E4002B", "#F5B3BF"])
    else:
        vals = sorted((rng.randint(20, 200) for _ in COUNTRIES), reverse=True)
        ax.barh(COUNTRIES[::-1], vals[::-1], color="#E4002B")
        ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    return buf.getvalue()


def textbox(slide, left, top, width, height, text, size=14, bold=False, color=INK, name=None):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if name:
        box.name = name
    return box


def build_deck() -> tuple[dict, dict]:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    chart_tags: dict[str, list[str]] = {}
    kpi_tags: dict[str, list[list[str]]] = {}

    def slide(title: str):
        s = prs.slides.add_slide(blank)
        textbox(s, 0.6, 0.35, 12, 0.8, title, size=28, bold=True)
        return s

    def cards(s, n, items):
        """KPI cards: a static caption plus a value box that gets tagged. Returns tag pairs."""
        pairs = []
        for i, (label, shown, key) in enumerate(items):
            x = 0.6 + i * (12.1 / n)
            textbox(s, x, 1.25, 3.8, 0.4, label, size=12, color=MUTED)
            textbox(s, x, 1.6, 3.8, 0.9, shown, size=36, bold=True, color=ACCENT)
            pairs.append([shown, key])
        return pairs

    def pictures(s, keys_kinds):
        for i, (_, kind) in enumerate(keys_kinds):
            s.shapes.add_picture(io.BytesIO(chart_png(kind, i + len(chart_tags))),
                                 Inches(0.6 + i * 6.3), Inches(2.9), width=Inches(6.0))

    # 1 cover
    s = prs.slides.add_slide(blank)
    textbox(s, 0.8, 2.4, 11.5, 1.2, "Analytics KPI Overview", size=44, bold=True)
    textbox(s, 0.8, 3.7, 11.5, 0.6, "Fictional demo deck — Contoso Group", size=18, color=MUTED)
    textbox(s, 0.8, 4.4, 6, 0.6, "June 2026 Data", size=20, color=ACCENT)
    kpi_tags["1"] = [["June 2026 Data", "report_month_label"]]

    # 2 academy
    s = slide("Analytics Academy")
    kpi_tags["2"] = cards(s, 3, [("Total learners", "100", "academy_total_learners"),
                                 ("Countries", "5", "academy_countries"),
                                 ("Companies", "9", "academy_companies")])
    chart_tags["2"] = ["academy_learners_by_country", "academy_learners_by_language"]
    pictures(s, [("a", "bar"), ("b", "donut")])

    # 3 trainings & guild
    s = slide("Databricks Trainings and Lakehouse Guild")
    kpi_tags["3"] = cards(s, 3, [("Training participants", "110", "db_participants"),
                                 ("Guild members", "120", "db_guild_members"),
                                 ("Training coverage", "0.30%", "db_training_coverage")])
    chart_tags["3"] = ["db_training_participation_country", "db_guild_participation_country"]
    pictures(s, [("a", "bar"), ("b", "bar")])

    # 4 community
    s = slide("Analytics Community")
    kpi_tags["4"] = cards(s, 3, [("Total registrations", "130", "community_total_registrations"),
                                 ("Teams members", "140", "community_teams_members"),
                                 ("Countries", "6", "community_countries_represented")])
    chart_tags["4"] = ["community_top_entities", "community_geographic_footprint"]
    pictures(s, [("a", "bar"), ("b", "bar")])

    # 5 events
    s = slide("Community Events")
    kpi_tags["5"] = cards(s, 3, [("Events", "2", "events_count"),
                                 ("Total participants", "150", "events_total_participants"),
                                 ("Unique participants", "160", "events_unique_participants")])
    chart_tags["5"] = ["events_participants_per_event", "events_participation_country"]
    pictures(s, [("a", "bar"), ("b", "bar")])

    # 6 manually maintained slide — listed in config/deck_layout.json, so the refresh never touches it
    s = slide("Notes (maintained by hand)")
    textbox(s, 0.6, 1.4, 11.5, 0.6, "This slide is listed under manual_slides in config/deck_layout.json,",
            size=16, color=MUTED)
    textbox(s, 0.6, 1.9, 11.5, 0.6, "so even its tagged shapes are never refreshed.", size=16, color=MUTED)
    textbox(s, 0.6, 2.8, 6, 0.9, "Edited by hand", size=32, bold=True, color=ACCENT)
    kpi_tags["6"] = [["Edited by hand", "group_employees"]]

    untagged = ROOT / "demo_data" / "Analytics_KPIs_demo_untagged.pptx"
    prs.save(untagged)
    spec = {"chart_tags": chart_tags, "kpi_tags": kpi_tags}
    return spec, {"untagged": untagged}


def main() -> None:
    build_excel()
    spec, paths = build_deck()
    spec_path = ROOT / "demo_data" / "demo_template_spec.json"
    spec_path.write_text(json.dumps(spec, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    out = ROOT / "templates" / "Analytics_KPIs_demo_template_tagged.pptx"
    kpis, charts = template_tagger.tag_template(paths["untagged"], out, spec_path)
    print(f"Excel files: {len(list(EXCEL.glob('*.xlsx')))} | tagged {kpis} KPI boxes, {charts} charts -> {out.name}")


if __name__ == "__main__":
    main()
